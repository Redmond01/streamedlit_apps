from __future__ import annotations

import re
from typing import Any
import openpyxl

from core.memory import spool_uploaded_file, cleanup_files, force_gc
from .common import cell_date_to_text, choose_block, normalize_pos, parse_money

POS_TYPES = ("SGR", "OBX", "AGS", "TA")
MIN_NUMBER_TYPES = {"SGR", "OBX"}


def find_cashbook_blocks(ws: Any) -> list[dict[str, Any]]:
    """Locates columns in row 4 containing 'cashbook'."""
    blocks = []
    max_c = ws.max_column or 50
    for col in range(1, max_c + 1):
        value = ws.cell(4, col).value
        if isinstance(value, str) and "cashbook" in value.lower():
            terminal_col = col - 1
            date = ""
            for date_col in range(max(1, terminal_col), min(max_c, col + 2) + 1):
                cell_val = ws.cell(1, date_col).value
                date = cell_date_to_text(cell_val) or date
            blocks.append({
                "date": date,
                "terminal_col": terminal_col,
                "cashbook_col": col,
            })
    return blocks


def parse_pos(value: Any) -> tuple[str, int] | None:
    text = str(value or "").strip().upper()
    match = re.match(r"^(SGR|OBX|AGS|T\.?\s*A|TA)\s*(\d+)\b", text)
    if not match:
        return None
    pos_type = match.group(1).replace(".", "").replace(" ", "")
    return pos_type, int(match.group(2))


def amount_value(value: Any) -> float | int | None:
    if value in (None, ""):
        return None
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return None
    if amount == 0:
        return None
    return int(amount) if amount.is_integer() else amount


def extract_cashbook_loop(
    workbook_input: Any,
    date: str | None = None,
    block_number: int | None = None,
    min_number: int = 11,
    target_day: int | None = None,
    include_audit: bool = False,
) -> dict[str, Any]:
    """Extracts active POS terminal numbers grouped by terminal type for Chrome extension."""
    tmp_path = spool_uploaded_file(workbook_input, prefix="cashbook_")
    wb = None
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True, read_only=False)
        ws = wb.active
        blocks = find_cashbook_blocks(ws)
        if not blocks:
            raise ValueError("No CASHBOOK EOD blocks found in row 4 of the active sheet.")

        block = choose_block(blocks, date=date, block_number=block_number)

        groups: dict[str, list[int]] = {pos_type: [] for pos_type in POS_TYPES}
        skipped_first_ten: list[dict[str, Any]] = []
        ignored_other: list[dict[str, Any]] = []

        max_r = ws.max_row or 200
        for row in range(5, max_r + 1):
            terminal = ws.cell(row, block["terminal_col"]).value
            amount = amount_value(ws.cell(row, block["cashbook_col"]).value)
            if terminal is None or amount is None:
                continue

            parsed = parse_pos(terminal)
            if not parsed:
                ignored_other.append({"row": row, "terminal": str(terminal).strip().upper(), "cashbookEod": amount})
                continue

            pos_type, number = parsed
            label = f"{pos_type} {number}"

            if pos_type in MIN_NUMBER_TYPES and number < min_number:
                skipped_first_ten.append({"row": row, "terminal": label, "cashbookEod": amount})
                continue

            if number not in groups[pos_type]:
                groups[pos_type].append(number)

        for key in groups:
            groups[key].sort()

        parsed_target_day = target_day
        if parsed_target_day is None and block["date"]:
            try:
                parsed_target_day = int(block["date"][-2:])
            except ValueError:
                parsed_target_day = None

        result: dict[str, Any] = {
            "date": block["date"],
            "targetDay": parsed_target_day,
            "minNumber": min_number,
            "groups": groups,
        }

        if include_audit:
            result["skippedFirstTen"] = skipped_first_ten
            result["ignoredOther"] = ignored_other

        return result
    finally:
        if wb:
            try:
                wb.close()
            except Exception:
                pass
        cleanup_files(tmp_path)
        force_gc()
