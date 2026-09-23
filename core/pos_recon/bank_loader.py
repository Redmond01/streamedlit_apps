from __future__ import annotations

import json
from pathlib import Path
import re
import tempfile
from typing import Any

import openpyxl
from openpyxl.comments import Comment

from core.memory import cleanup_files, force_gc, spool_uploaded_file
from .common import choose_block, find_report_blocks, normalize_pos, parse_money


def parse_records(payload: Any) -> list[dict[str, Any]]:
    """Extracts records list from raw JSON payload (list or wrapped dict)."""
    if isinstance(payload, list):
        return payload
    if isinstance(payload, dict):
        for key in ("records", "results", "rows", "data", "exportedRecords", "savedRecords"):
            if isinstance(payload.get(key), list):
                return payload[key]
    raise ValueError("JSON must be an exported records array or contain a records/results/rows list.")


def infer_date(records: list[dict[str, Any]]) -> str:
    """Attempts to infer the report date from records metadata."""
    for record in records:
        run_date = record.get("runDate")
        if run_date:
            return str(run_date)[:10]
    for record in records:
        date_text = str(record.get("date") or "")
        match = re.search(r"(\d{1,2})\s+[A-Za-z]{3},\s*(\d{4})", date_text)
        if match:
            return ""
    return ""


def build_record_map(records: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], int]:
    """Maps normalized POS labels to their respective record objects."""
    by_pos: dict[str, dict[str, Any]] = {}
    duplicate_count = 0
    for record in records:
        pos = normalize_pos(record.get("pos"))
        if not pos:
            continue
        if pos in by_pos:
            duplicate_count += 1
        by_pos[pos] = record
    return by_pos, duplicate_count


def load_json_to_workbook(
    workbook_input: Any,
    json_input: Any,
    date: str | None = None,
    block_number: int | None = None,
    failed_text: str = "FAILED",
) -> tuple[Path, dict[str, Any]]:
    """Fills matching POS rows in the reconciliation workbook's BANK STATEMENT column."""
    # Parse JSON content
    if hasattr(json_input, "read"):
        content = json_input.read()
        if isinstance(content, bytes):
            content = content.decode("utf-8")
        raw_json = json.loads(content)
    elif isinstance(json_input, (str, Path)):
        with open(json_input, "r", encoding="utf-8") as f:
            raw_json = json.load(f)
    else:
        raw_json = json_input

    records = parse_records(raw_json)
    if not date:
        date = infer_date(records) or None

    tmp_path = spool_uploaded_file(workbook_input, prefix="bank_wb_")
    wb = None
    output_path: Path | None = None

    try:
        wb = openpyxl.load_workbook(tmp_path)
        ws = wb.active
        blocks = find_report_blocks(ws)
        if not blocks:
            raise ValueError("No BANK STATEMENT report blocks found in the workbook.")

        block = choose_block(blocks, date=date, block_number=block_number)
        record_map, duplicate_count = build_record_map(records)

        matched = 0
        ok_written = 0
        status_written = 0
        workbook_pos_rows = 0
        cashbook_rows = 0
        cashbook_rows_missing_in_json = 0
        mapping_details = []

        max_r = ws.max_row or 200
        for row in range(5, max_r + 1):
            raw_terminal = ws.cell(row, block["terminal_col"]).value
            pos = normalize_pos(raw_terminal)
            if not pos:
                continue

            workbook_pos_rows += 1
            cashbook_value = ws.cell(row, block["cashbook_col"]).value
            has_cashbook_value = cashbook_value not in (None, "", 0)
            if has_cashbook_value:
                cashbook_rows += 1

            record = record_map.get(pos)
            if not record:
                if has_cashbook_value:
                    cashbook_rows_missing_in_json += 1
                continue

            matched += 1
            cell = ws.cell(row, block["bank_col"])
            status = str(record.get("status") or "").upper()

            if status == "OK":
                amount = parse_money(record.get("credit"))
                if amount is None:
                    cell.value = failed_text
                    status_written += 1
                else:
                    cell.value = amount
                    cell.number_format = "#,##0.00"
                    ok_written += 1
            else:
                cell.value = status or failed_text
                status_written += 1

            error = record.get("error") or ""
            saved_at = record.get("savedAt") or ""
            if error or status != "OK":
                cell.comment = Comment(
                    f"status={status or failed_text}\nerror={error}\nsavedAt={saved_at}",
                    "POSBOT",
                )

            mapping_details.append({
                "row": row,
                "terminal": pos,
                "status": status or failed_text,
                "credit": record.get("credit"),
                "has_comment": bool(error or status != "OK"),
            })

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", prefix="BANK_LOADED_") as out_tmp:
            output_path = Path(out_tmp.name)
        wb.save(output_path)

        stats = {
            "workbookDate": block["date"],
            "jsonRecords": len(records),
            "jsonUniquePos": len(record_map),
            "workbookPosRows": workbook_pos_rows,
            "cashbookRows": cashbook_rows,
            "matched": matched,
            "okWritten": ok_written,
            "statusWritten": status_written,
            "cashbookRowsMissingInJson": cashbook_rows_missing_in_json,
            "duplicatesInJson": duplicate_count,
            "mappingDetails": mapping_details,
        }

        return output_path, stats

    finally:
        if wb:
            try:
                wb.close()
            except Exception:
                pass
        cleanup_files(tmp_path)
        force_gc()
