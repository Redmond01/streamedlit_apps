from __future__ import annotations

from datetime import datetime
from pathlib import Path
import tempfile
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill

from core.memory import cleanup_files, force_gc, spool_uploaded_file
from .common import (
    cell_date_to_text,
    choose_block,
    find_report_blocks,
    normalize_pos,
    parse_money,
)


def daily_pos_sheet_name(date_str: str) -> str:
    """Returns sheet name formatted as Daily_POS_YYYY_MM_DD."""
    try:
        parsed = datetime.fromisoformat(date_str)
    except ValueError as error:
        raise ValueError("Date must use YYYY-MM-DD format (e.g. 2026-07-26).") from error
    return f"Daily_POS_{parsed:%Y_%m_%d}"


def choose_source_sheet(wb: Any, date_str: str, source_sheet: str | None = None) -> tuple[Any, str]:
    expected = source_sheet or daily_pos_sheet_name(date_str)
    if expected in wb.sheetnames:
        return wb[expected], expected

    available = ", ".join(wb.sheetnames)
    if source_sheet:
        raise ValueError(f"Source workbook does not contain sheet '{source_sheet}'. Available: {available}")

    raise ValueError(
        f"Source workbook does not contain expected sheet '{expected}' for date {date_str}. "
        f"Provide a sheet name override if needed. Available sheets: {available}"
    )


def find_header(headers: list[Any], wanted: str) -> int:
    wanted_clean = wanted.strip().lower()
    for index, value in enumerate(headers, start=1):
        if str(value or "").strip().lower() == wanted_clean:
            return index
    raise ValueError(f"Missing required source header: '{wanted}'")


def find_total_row(ws: Any, terminal_col: int) -> int:
    max_r = ws.max_row or 200
    for row in range(5, max_r + 1):
        if str(ws.cell(row, terminal_col).value or "").strip().upper() == "TOTAL":
            return row
    return max_r


def extract_eod_source(source_path: Path, date_str: str, source_sheet: str | None = None) -> dict[str, Any]:
    wb = openpyxl.load_workbook(source_path, data_only=True)
    try:
        ws, selected_sheet = choose_source_sheet(wb, date_str, source_sheet=source_sheet)

        headers = [cell.value for cell in ws[1]]
        timestamp_col = find_header(headers, "Timestamp")
        submission_id_col = find_header(headers, "Submission ID")
        station_col = find_header(headers, "Station")
        report_date_col = find_header(headers, "Report Date")
        full_pos_col = find_header(headers, "Full POS ID")
        amount_col = find_header(headers, "Amount")

        amount_by_pos: dict[str, float] = {}
        kept_rows = []
        skipped_wrong_date = 0
        skipped_invalid_terminal = []
        skipped_invalid_amount = []
        source_rows = 0

        max_r = ws.max_row or 500
        for row in range(2, max_r + 1):
            row_values = [ws.cell(row, col).value for col in range(1, len(headers) + 1)]
            if all(value in (None, "") for value in row_values):
                continue
            source_rows += 1

            report_date = cell_date_to_text(ws.cell(row, report_date_col).value)
            if report_date != date_str:
                skipped_wrong_date += 1
                continue

            pos = normalize_pos(ws.cell(row, full_pos_col).value)
            if not pos:
                skipped_invalid_terminal.append({
                    "row": row,
                    "fullPosId": ws.cell(row, full_pos_col).value,
                    "station": ws.cell(row, station_col).value,
                })
                continue

            amount = parse_money(ws.cell(row, amount_col).value)
            if amount is None:
                skipped_invalid_amount.append({
                    "row": row,
                    "pos": pos,
                    "amount": ws.cell(row, amount_col).value,
                })
                continue

            amount_by_pos[pos] = round(amount_by_pos.get(pos, 0.0) + amount, 2)
            kept_rows.append({
                "sourceRow": row,
                "timestamp": cell_date_to_text(ws.cell(row, timestamp_col).value),
                "submissionId": ws.cell(row, submission_id_col).value,
                "station": ws.cell(row, station_col).value,
                "sourcePos": ws.cell(row, full_pos_col).value,
                "pos": pos,
                "amount": amount,
            })

        return {
            "sourceRows": source_rows,
            "keptRows": kept_rows,
            "skippedWrongDate": skipped_wrong_date,
            "skippedInvalidTerminal": skipped_invalid_terminal,
            "skippedInvalidAmount": skipped_invalid_amount,
            "amountByPos": amount_by_pos,
            "sourceSheet": selected_sheet,
        }
    finally:
        wb.close()


def write_eod_to_workbook(
    source_input: Any,
    target_workbook_input: Any,
    date_str: str,
    source_sheet: str | None = None,
    block_number: int | None = None,
) -> tuple[Path, dict[str, Any]]:
    """Syncs Moniepoint EOD transactions into reconciliation workbook's CASHBOOK column."""
    if not date_str:
        raise ValueError("Report Date is required to filter EOD records safely.")

    src_tmp = spool_uploaded_file(source_input, prefix="eod_src_")
    tgt_tmp = spool_uploaded_file(target_workbook_input, prefix="recon_tgt_")

    target_wb = None
    output_path: Path | None = None

    try:
        # Extract EOD data from source file
        extracted = extract_eod_source(src_tmp, date_str, source_sheet=source_sheet)

        # Open target workbook for modification
        target_wb = openpyxl.load_workbook(tgt_tmp)
        ws = target_wb.active
        blocks = find_report_blocks(ws)
        if not blocks:
            raise ValueError("No BANK STATEMENT report blocks found in target workbook.")

        block = choose_block(blocks, date=date_str, block_number=block_number)

        total_row = find_total_row(ws, block["terminal_col"])
        data_end_row = total_row - 1
        unmatched = set(extracted["amountByPos"])
        mapped_rows = []

        try:
            ws.cell(1, block["bank_col"]).value = datetime.fromisoformat(date_str)
            ws.cell(1, block["bank_col"]).number_format = "yyyy-mm-dd"
        except Exception:
            ws.cell(1, block["bank_col"]).value = date_str

        for row in range(5, data_end_row + 1):
            ws.cell(row, block["cashbook_col"]).value = None
            ws.cell(row, block["bank_col"]).value = None
            ws.cell(row, block["difference_col"]).value = (
                f"=SUM({ws.cell(row, block['cashbook_col']).coordinate}-"
                f"{ws.cell(row, block['bank_col']).coordinate})"
            )

            pos = normalize_pos(ws.cell(row, block["terminal_col"]).value)
            if not pos or pos not in extracted["amountByPos"]:
                continue

            amount = extracted["amountByPos"][pos]
            ws.cell(row, block["cashbook_col"]).value = amount
            ws.cell(row, block["cashbook_col"]).number_format = "#,##0.00"
            mapped_rows.append({
                "row": row,
                "workbookPos": ws.cell(row, block["terminal_col"]).value,
                "pos": pos,
                "amount": amount,
            })
            unmatched.discard(pos)

        ws.cell(total_row, block["cashbook_col"]).value = f"=SUM({ws.cell(5, block['cashbook_col']).coordinate}:{ws.cell(data_end_row, block['cashbook_col']).coordinate})"
        ws.cell(total_row, block["bank_col"]).value = f"=SUM({ws.cell(5, block['bank_col']).coordinate}:{ws.cell(data_end_row, block['bank_col']).coordinate})"
        ws.cell(total_row, block["difference_col"]).value = (
            f"=SUM({ws.cell(total_row, block['cashbook_col']).coordinate}-"
            f"{ws.cell(total_row, block['bank_col']).coordinate})"
        )

        _write_audit_sheet(
            target_wb,
            source_filename=getattr(source_input, "name", "EOD_Source.xlsx"),
            target_filename=getattr(target_workbook_input, "name", "Reconciliation.xlsx"),
            date_str=date_str,
            extracted=extracted,
            mapped_rows=mapped_rows,
            unmatched=sorted(unmatched),
        )

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", prefix=f"EOD_SYNC_{date_str.replace('-', '')}_") as out_tmp:
            output_path = Path(out_tmp.name)
        target_wb.save(output_path)

        result_stats = {
            "sourceSheet": extracted["sourceSheet"],
            "sourceRows": extracted["sourceRows"],
            "keptRows": len(extracted["keptRows"]),
            "skippedWrongDate": extracted["skippedWrongDate"],
            "skippedInvalidTerminal": len(extracted["skippedInvalidTerminal"]),
            "skippedInvalidAmount": len(extracted["skippedInvalidAmount"]),
            "uniqueTerminals": len(extracted["amountByPos"]),
            "mappedRows": len(mapped_rows),
            "unmatchedSourceTerminals": sorted(unmatched),
        }
        return output_path, result_stats

    finally:
        if target_wb:
            try:
                target_wb.close()
            except Exception:
                pass
        cleanup_files(src_tmp, tgt_tmp)
        force_gc()


def _write_audit_sheet(
    wb: Any,
    source_filename: str,
    target_filename: str,
    date_str: str,
    extracted: dict[str, Any],
    mapped_rows: list[dict[str, Any]],
    unmatched: list[str],
) -> None:
    audit_sheet_name = "EOD Extract Audit"
    if audit_sheet_name in wb.sheetnames:
        del wb[audit_sheet_name]
    ws = wb.create_sheet(audit_sheet_name)

    rows: list[list[Any]] = [
        ["Metric", "Value"],
        ["Source Workbook", source_filename],
        ["Source Sheet", extracted.get("sourceSheet", "")],
        ["Target Workbook", target_filename],
        ["Target Date", date_str],
        ["Mapped Records", len(mapped_rows)],
        ["Unmatched Source Terminals", len(unmatched)],
        ["Rows Kept for Target Date", len(extracted["keptRows"])],
        ["Rows Skipped for Other Dates", extracted["skippedWrongDate"]],
        [],
        ["Mapped Row", "Workbook Terminal", "Normalized Terminal", "EOD Amount"],
    ]
    rows.extend([r["row"], r["workbookPos"], r["pos"], r["amount"]] for r in mapped_rows)
    rows.extend([[], ["Unmatched Source Terminal"]])
    rows.extend([terminal] for terminal in unmatched)
    rows.extend([[], ["Kept Source Row", "Timestamp", "Submission ID", "Station", "Source POS", "Normalized POS", "Amount"]])
    rows.extend([
        r["sourceRow"],
        r["timestamp"],
        r["submissionId"],
        r["station"],
        r["sourcePos"],
        r["pos"],
        r["amount"],
    ] for r in extracted["keptRows"])

    for row_idx, row_values in enumerate(rows, start=1):
        for col_idx, val in enumerate(row_values, start=1):
            ws.cell(row_idx, col_idx).value = val

    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in ws[10]:
        cell.font = header_font
        cell.fill = PatternFill(fill_type="solid", fgColor="5B9BD5")

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_length + 2, 12), 45)
