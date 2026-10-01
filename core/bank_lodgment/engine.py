from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
import logging
from pathlib import Path
import tempfile
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from core.gear.utils import decimal_to_excel
from core.memory import cleanup_files, force_gc, spool_uploaded_file
from .config import BankLodgmentConfig
from .matcher import format_lodgment_time, normalize_bank_name, resolve_station_for_record
from .parser import BankLodgmentRecord, normalize_iso_to_sheet_date, parse_lodgment_json

logger = logging.getLogger(__name__)

GREEN_FILL_VIBRANT = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GREEN_FONT_DARK = Font(name="Calibri", size=11, color="006100", bold=True)
GREEN_FILL_ROW = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

RED_FILL_VIBRANT = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT_DARK = Font(name="Calibri", size=11, color="9C0006", bold=True)
RED_FILL_ROW = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
AMOUNT_FORMAT = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'


def apply_lodgment_sheet_styles(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    color_mode: str = "row",
    total_row: int | None = None,
) -> None:
    """Applies green/red fills and accounting number formats to lodgment rows."""
    max_r = total_row if total_row is not None else ws.max_row
    for r in range(2, max_r + 1):
        st_val = ws.cell(r, 1).value
        if not st_val or str(st_val).strip().upper() == "TOTAL":
            break

        # Ensure amount column has accounting number format
        amt_cell = ws.cell(r, 3)
        if amt_cell.value is not None:
            amt_cell.number_format = AMOUNT_FORMAT

        conf_val = str(ws.cell(r, 7).value or "").strip().upper()
        if conf_val == "YES":
            for c in (7, 8, 9):
                ws.cell(r, c).fill = GREEN_FILL_VIBRANT
                ws.cell(r, c).font = GREEN_FONT_DARK
                ws.cell(r, c).alignment = CENTER_ALIGN
            if color_mode == "row":
                for c in range(1, 7):
                    ws.cell(r, c).fill = GREEN_FILL_ROW
        elif conf_val == "NO":
            ws.cell(r, 7).fill = RED_FILL_VIBRANT
            ws.cell(r, 7).font = RED_FONT_DARK
            ws.cell(r, 7).alignment = CENTER_ALIGN
            if color_mode == "row":
                for c in range(1, 7):
                    ws.cell(r, c).fill = RED_FILL_ROW


@dataclass
class BankLodgmentResult:
    output_path: Path | None
    target_sheet_name: str
    total_records: int
    confirmed_count: int
    duplicate_count: int
    unmatched_count: int
    total_confirmed_amount: Decimal
    union_confirmed_amount: Decimal
    fcmb_confirmed_amount: Decimal
    logs: list[dict[str, Any]] = field(default_factory=list)


def process_bank_lodgment(
    master_file: Any,
    json_file: Any,
    config: BankLodgmentConfig,
    dry_run: bool = False,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> BankLodgmentResult:
    """Processes bank deposit JSON dump against master lodgment workbook."""
    config.validate()

    master_path = spool_uploaded_file(master_file, prefix="lodgment_master_")
    json_path = spool_uploaded_file(json_file, prefix="lodgment_json_")
    output_path: Path | None = None
    master_wb = None

    logs: list[dict[str, Any]] = []
    confirmed_count = 0
    duplicate_count = 0
    unmatched_count = 0

    total_confirmed_amount = Decimal("0")
    union_confirmed_amount = Decimal("0")
    fcmb_confirmed_amount = Decimal("0")

    try:
        # 1. Parse JSON
        meta, records = parse_lodgment_json(json_path)

        # 2. Determine target sheet date
        raw_target_date = (
            config.target_date
            or meta.get("target_reconciliation_date")
            or meta.get("date_of_reconciliation_pulled")
        )
        if not raw_target_date:
            raise ValueError("Target reconciliation date could not be determined from configuration or JSON.")

        target_sheet_name = normalize_iso_to_sheet_date(raw_target_date)

        # 3. Load Master Workbook
        master_wb = openpyxl.load_workbook(master_path, data_only=False)

        # 4. Sheet resolution or cloning
        if target_sheet_name not in master_wb.sheetnames:
            if not config.auto_create_sheet:
                available = ", ".join(master_wb.sheetnames)
                raise ValueError(
                    f"Sheet '{target_sheet_name}' not found in workbook. Available sheets: {available}"
                )
            # Find template or most recent date sheet to clone
            date_sheets = [s for s in master_wb.sheetnames if s != "SUMMARY"]
            if not date_sheets:
                raise ValueError("No date worksheets available to clone template from.")
            source_template_name = date_sheets[-1]
            source_ws = master_wb[source_template_name]
            target_ws = master_wb.copy_worksheet(source_ws)
            target_ws.title = target_sheet_name

            # Reset day-specific columns (Amount, Trans. Ref, Time, Confirmed Y/N, Confirmed By, Flag)
            for r in range(2, target_ws.max_row + 1):
                st_val = target_ws.cell(r, 1).value
                if str(st_val).strip().upper() == "TOTAL":
                    break
                target_ws.cell(r, 3).value = None   # Amount
                target_ws.cell(r, 5).value = None   # Trans. Ref
                target_ws.cell(r, 6).value = None   # Lodgement Time
                target_ws.cell(r, 7).value = "NO"   # Confirmed Y/N
                target_ws.cell(r, 8).value = None   # Confirmed By
                target_ws.cell(r, 9).value = None   # Flag

            # Link to SUMMARY sheet if exists
            if "SUMMARY" in master_wb.sheetnames:
                summary_ws = master_wb["SUMMARY"]
                next_row = summary_ws.max_row + 1
                try:
                    dt_obj = datetime.strptime(target_sheet_name, "%d-%m-%Y")
                    summary_ws.cell(next_row, 1).value = dt_obj
                    summary_ws.cell(next_row, 2).value = f"=SUMIFS('{target_sheet_name}'!C:C,'{target_sheet_name}'!B:B,\"Zenith\")"
                    summary_ws.cell(next_row, 3).value = f"=SUMIFS('{target_sheet_name}'!C:C,'{target_sheet_name}'!B:B,\"FCMB\")"
                    summary_ws.cell(next_row, 4).value = f"=SUMIFS('{target_sheet_name}'!C:C,'{target_sheet_name}'!B:B,\"Union\")"
                    summary_ws.cell(next_row, 5).value = f"=SUM(B{next_row}:D{next_row})"
                except Exception as exc:
                    logger.warning("Could not append row to SUMMARY sheet: %s", exc)
        else:
            target_ws = master_wb[target_sheet_name]

        # 5. Index existing stations in target worksheet
        station_to_rows: dict[str, list[int]] = {}
        depositors_map: dict[str, str] = {}
        total_row_idx = None

        for r in range(2, target_ws.max_row + 1):
            st_val = target_ws.cell(r, 1).value
            if not st_val:
                continue
            st_clean = str(st_val).strip()
            if st_clean.upper() == "TOTAL":
                total_row_idx = r
                break
            station_to_rows.setdefault(st_clean, []).append(r)
            dep_val = target_ws.cell(r, 4).value
            if dep_val:
                depositors_map[str(dep_val).strip()] = st_clean

        if total_row_idx is None:
            total_row_idx = target_ws.max_row + 1

        total_recs = len(records)

        # 6. Iterate and match records
        for idx, rec in enumerate(records):
            if progress_callback:
                progress_callback(idx + 1, total_recs, rec.message_id or f"Record {idx + 1}")

            # Check duplicates
            if rec.is_duplicate:
                duplicate_count += 1
                logs.append({
                    "time": rec.time,
                    "station": "-",
                    "amount": float(rec.amount),
                    "bank": rec.bank_officer_bank,
                    "officer": rec.bank_officer_name,
                    "manager": rec.manager_handle or rec.manager_phone,
                    "target_row": "-",
                    "status": "Skipped (Duplicate)",
                    "message": "Exact duplicate confirmation message in export",
                })
                continue

            # Resolve station
            resolved_station, strategy, score = resolve_station_for_record(
                record=rec,
                config=config,
                known_stations=list(station_to_rows.keys()),
                depositors_map=depositors_map,
            )

            if not resolved_station or resolved_station not in station_to_rows:
                unmatched_count += 1
                logs.append({
                    "time": rec.time,
                    "station": resolved_station or "UNMATCHED",
                    "amount": float(rec.amount),
                    "bank": rec.bank_officer_bank,
                    "officer": rec.bank_officer_name,
                    "manager": rec.manager_handle or rec.manager_phone,
                    "target_row": "-",
                    "status": "Unmatched",
                    "message": f"Could not match to worksheet stations (strategy={strategy}, score={score:.2f})",
                })
                continue

            # Find matching row for this station
            candidate_rows = station_to_rows[resolved_station]
            target_row = None

            # 1. Prefer unconfirmed row or row with exact matching amount
            for r in candidate_rows:
                conf_val = str(target_ws.cell(r, 7).value or "").strip().upper()
                amt_val = target_ws.cell(r, 3).value
                if amt_val is not None and Decimal(str(amt_val)) == rec.amount:
                    target_row = r
                    break
                if conf_val != "YES" and amt_val in (None, ""):
                    target_row = r
                    break

            # 2. If no empty or matching row, pick first unconfirmed
            if target_row is None:
                for r in candidate_rows:
                    conf_val = str(target_ws.cell(r, 7).value or "").strip().upper()
                    if conf_val != "YES":
                        target_row = r
                        break

            # 3. If all rows already confirmed, duplicate row right below
            if target_row is None:
                last_r = candidate_rows[-1]
                target_row = last_r + 1
                if not dry_run:
                    target_ws.insert_rows(target_row)
                    # Shift all recorded rows down by 1
                    for st_name, rows_list in station_to_rows.items():
                        station_to_rows[st_name] = [
                            row + 1 if row >= target_row else row for row in rows_list
                        ]
                    station_to_rows[resolved_station].append(target_row)
                    total_row_idx += 1

                    # Copy station and bank name
                    target_ws.cell(target_row, 1).value = resolved_station
                    target_ws.cell(target_row, 2).value = target_ws.cell(last_r, 2).value

            # Format fields to write
            norm_bank = normalize_bank_name(rec.bank_officer_bank)
            formatted_time = format_lodgment_time(rec.time)

            if not dry_run:
                target_ws.cell(target_row, 3).value = decimal_to_excel(rec.amount)
                if not target_ws.cell(target_row, 6).value:
                    target_ws.cell(target_row, 6).value = formatted_time
                target_ws.cell(target_row, 7).value = "YES"
                target_ws.cell(target_row, 8).value = rec.bank_officer_name.upper()
                target_ws.cell(target_row, 9).value = config.flag_value

            confirmed_count += 1
            total_confirmed_amount += rec.amount
            if "UNION" in norm_bank:
                union_confirmed_amount += rec.amount
            elif "FCMB" in norm_bank or "ACCESS" in norm_bank:
                fcmb_confirmed_amount += rec.amount

            logs.append({
                "time": formatted_time,
                "station": resolved_station,
                "amount": float(rec.amount),
                "bank": norm_bank,
                "officer": rec.bank_officer_name.upper(),
                "manager": rec.manager_handle or rec.manager_phone,
                "target_row": f"Row {target_row}",
                "status": "Preview" if dry_run else "Confirmed",
                "message": f"Matched via {strategy} (conf={score:.2f})",
            })

        # Ensure styling and bottom TOTAL formula are accurate
        if not dry_run:
            apply_lodgment_sheet_styles(target_ws, color_mode=config.color_mode, total_row=total_row_idx)
            target_ws.cell(total_row_idx, 3).value = f"=SUM(C2:C{total_row_idx - 1})"
            target_ws.cell(total_row_idx, 3).number_format = AMOUNT_FORMAT
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", prefix="BANK_LODGEMENT_") as out_tmp:
                output_path = Path(out_tmp.name)
            master_wb.save(output_path)
            master_wb.close()
            master_wb = None

    finally:
        if master_wb:
            try:
                master_wb.close()
            except Exception:
                pass
        cleanup_files(master_path, json_path)
        force_gc()

    return BankLodgmentResult(
        output_path=output_path,
        target_sheet_name=target_sheet_name,
        total_records=len(records) if "records" in locals() else 0,
        confirmed_count=confirmed_count,
        duplicate_count=duplicate_count,
        unmatched_count=unmatched_count,
        total_confirmed_amount=total_confirmed_amount,
        union_confirmed_amount=union_confirmed_amount,
        fcmb_confirmed_amount=fcmb_confirmed_amount,
        logs=logs,
    )
