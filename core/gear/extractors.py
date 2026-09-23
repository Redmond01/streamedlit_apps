from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
import re
from typing import Any, NamedTuple

from openpyxl.worksheet.worksheet import Worksheet

from .utils import is_blank, normalize_header, normalize_tank_name, parse_number


REQUIRED_HEADERS = {
    "TANK": ("TANK",),
    "SALES": ("SALES",),
    "ACTUAL GEAR": ("ACTUAL GEAR", "GEAR"),
    "EXPCTD GEAR": ("EXPCTD GEAR", "EXPECTED GEAR"),
}

TOTAL_PATTERN = re.compile(r"\bT[O0]TAL\s+(PMS|AGO)\s*\d*\b", re.IGNORECASE)


@dataclass(frozen=True)
class TankValues:
    tank: str
    dipping_sales: Decimal | None
    actual_gear: Decimal | None
    expected_gear: Decimal | None
    joined_or_empty: bool = False


class MeterTotal(NamedTuple):
    tank: str
    value: Decimal
    row: int
    column: int


def extract_tank_data(sheet: Worksheet) -> dict[str, TankValues]:
    """Extracts dipping sales, actual gear, and expected gear per tank."""
    header_row, columns = _find_tank_header(sheet)
    if header_row is None:
        return {}
    return _extract_tank_table(sheet, header_row, columns)


def _find_tank_header(sheet: Worksheet) -> tuple[int | None, dict[str, int]]:
    headers = _find_tank_headers(sheet)
    if not headers:
        return None, {}
    return sorted(headers, key=lambda item: (item[1]["TANK"], item[0]))[0]


def _find_tank_headers(sheet: Worksheet) -> list[tuple[int, dict[str, int]]]:
    headers: list[tuple[int, dict[str, int]]] = []
    max_r = sheet.max_row or 100
    for row in sheet.iter_rows(min_row=1, max_row=max_r):
        header_map = _map_headers([cell.value for cell in row])
        if all(key in header_map for key in REQUIRED_HEADERS):
            headers.append((row[0].row, header_map))
            continue
        if all(key in header_map for key in ("SALES", "ACTUAL GEAR", "EXPCTD GEAR")):
            dipping_col = _first_dipping_header_col([cell.value for cell in row])
            if dipping_col and dipping_col > 1:
                header_map["TANK"] = dipping_col - 1
                headers.append((row[0].row, header_map))
    return headers


def _extract_tank_table(sheet: Worksheet, header_row: int, columns: dict[str, int]) -> dict[str, TankValues]:
    tanks: dict[str, TankValues] = {}
    blank_run = 0
    max_r = sheet.max_row or (header_row + 50)
    for row_idx in range(header_row + 1, max_r + 1):
        tank_raw = sheet.cell(row_idx, columns["TANK"]).value
        if isinstance(tank_raw, str) and "TOTAL" in tank_raw.upper():
            continue
        tank = normalize_tank_name(tank_raw)
        if tank is None:
            blank_run += 1 if all(is_blank(sheet.cell(row_idx, col).value) for col in columns.values()) else 0
            if blank_run >= 8:
                break
            continue
        blank_run = 0
        if not (tank.startswith("PMS") or tank.startswith("AGO")):
            continue
        dipping_sales = parse_number(sheet.cell(row_idx, columns["SALES"]).value)
        actual_gear = parse_number(sheet.cell(row_idx, columns["ACTUAL GEAR"]).value)
        expected_gear = parse_number(sheet.cell(row_idx, columns["EXPCTD GEAR"]).value)
        joined_or_empty = all(is_blank(sheet.cell(row_idx, columns[key]).value) for key in ("SALES", "ACTUAL GEAR", "EXPCTD GEAR"))
        tanks[tank] = TankValues(tank, dipping_sales, actual_gear, expected_gear, joined_or_empty)
    return tanks


def _map_headers(values: list[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for idx, value in enumerate(values, start=1):
        header = normalize_header(value)
        for canonical, aliases in REQUIRED_HEADERS.items():
            if header in aliases and canonical not in mapped:
                mapped[canonical] = idx
    return mapped


def _first_dipping_header_col(values: list[Any]) -> int | None:
    for idx, value in enumerate(values, start=1):
        header = normalize_header(value)
        if header in {"OPENNG DIPPG", "OPENING DIPPG", "OPENING DIPPING", "OPENNG DIPPING"}:
            return idx
    return None


def extract_meter_sales(sheet: Worksheet, tank_rows: dict[str, TankValues] | None = None) -> dict[str, Decimal]:
    """Extracts meter sales per tank by finding TOTAL PMS / TOTAL AGO cells."""
    entries: list[MeterTotal] = []
    max_r = sheet.max_row or 100
    for row in sheet.iter_rows(min_row=1, max_row=max_r):
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or not TOTAL_PATTERN.search(value):
                continue
            tank = normalize_tank_name(value)
            if tank is None:
                continue
            number = parse_number(sheet.cell(cell.row, cell.column + 1).value)
            if number is None:
                number = _first_number_below_value_cell(sheet, cell.row + 1, cell.column, cell.column + 1)
            if number is not None:
                _append_meter_total(entries, MeterTotal(tank, number, cell.row, cell.column))
    return _map_meter_totals(entries, tank_rows or {})


def _first_number_below_value_cell(sheet: Worksheet, start_row: int, label_col: int, value_col: int) -> Decimal | None:
    max_r = sheet.max_row or (start_row + 5)
    for row in range(start_row, min(max_r, start_row + 3) + 1):
        label_value = sheet.cell(row, label_col).value
        value = sheet.cell(row, value_col).value
        if _looks_like_new_total_or_tank(label_value) or _looks_like_new_total_or_tank(value):
            break
        number = parse_number(sheet.cell(row, value_col).value)
        if number is not None:
            return number
    return None


def _looks_like_new_total_or_tank(value: object) -> bool:
    if not isinstance(value, str):
        return False
    text = value.strip().upper()
    return bool(TOTAL_PATTERN.search(text) or normalize_tank_name(text))


def _map_meter_totals(entries: list[MeterTotal], tank_rows: dict[str, TankValues]) -> dict[str, Decimal]:
    totals: dict[str, Decimal] = {}
    for tank, entry in _first_section_entries(entries).items():
        totals[tank] = entry.value

    if not tank_rows:
        return totals

    for fuel in ("PMS", "AGO"):
        fuel_entries = [entry for entry in entries if entry.tank.startswith(fuel)]
        fuel_tank_rows = _sorted_fuel_tanks(tank_rows, fuel)
        if not fuel_entries or not fuel_tank_rows:
            continue

        joined_tanks = {tank for tank in fuel_tank_rows if tank_rows[tank].joined_or_empty}
        non_joined_tanks = [tank for tank in fuel_tank_rows if tank not in joined_tanks]

        if joined_tanks:
            _move_joined_totals_forward(totals, fuel_entries, fuel_tank_rows, joined_tanks, non_joined_tanks)

    return totals


def _append_meter_total(entries: list[MeterTotal], incoming: MeterTotal) -> None:
    for existing in entries:
        if existing.row == incoming.row and existing.tank == incoming.tank and abs(existing.value - incoming.value) <= Decimal("1"):
            return
    entries.append(incoming)


def _first_section_entries(entries: list[MeterTotal]) -> dict[str, MeterTotal]:
    selected: dict[str, MeterTotal] = {}
    for entry in entries:
        existing = selected.get(entry.tank)
        if existing is None or (entry.column, entry.row) < (existing.column, existing.row):
            selected[entry.tank] = entry
    return selected


def _move_joined_totals_forward(
    totals: dict[str, Decimal],
    entries: list[MeterTotal],
    fuel_tanks: list[str],
    joined_tanks: set[str],
    non_joined_tanks: list[str],
) -> None:
    for entry in entries:
        if entry.tank not in joined_tanks or entry.value == 0:
            continue
        target_tank = _next_non_joined_tank(entry.tank, fuel_tanks, non_joined_tanks)
        if target_tank is None:
            continue
        existing = totals.get(target_tank)
        totals[target_tank] = entry.value if existing is None or existing == 0 else existing + entry.value
        totals.pop(entry.tank, None)


def _next_non_joined_tank(tank: str, fuel_tanks: list[str], non_joined_tanks: list[str]) -> str | None:
    if tank not in fuel_tanks:
        return None
    start_index = fuel_tanks.index(tank) + 1
    for candidate in fuel_tanks[start_index:]:
        if candidate in non_joined_tanks:
            return candidate
    return None


def _sorted_fuel_tanks(tank_rows: dict[str, TankValues], fuel: str) -> list[str]:
    return sorted(
        [tank for tank in tank_rows if tank.startswith(fuel)],
        key=lambda value: int(re.search(r"\d+", value).group(0)),
    )
