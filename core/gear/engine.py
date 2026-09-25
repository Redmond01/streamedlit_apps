from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
import logging
from pathlib import Path
import tempfile
from typing import Any, Callable

import openpyxl
from openpyxl.utils import get_column_letter

from core.memory import cleanup_files, force_gc, spool_uploaded_file
from .config import GearConfig
from .extractors import TankValues, extract_meter_sales, extract_tank_data
from .formulas import build_average_formula, build_total_formulas
from .matcher import is_protected_column, match_station_columns
from .utils import decimal_to_excel
from .xml_preserver import preserve_gear_extensions

logger = logging.getLogger(__name__)


@dataclass
class GearBatchResult:
    output_path: Path | None
    processed_count: int
    skipped_count: int
    logs: list[dict[str, Any]]
    total_formulas_written: list[dict[str, str]]


def _is_zero_tank_record(tank_values: TankValues) -> bool:
    values = (tank_values.dipping_sales, tank_values.actual_gear, tank_values.expected_gear)
    return all(v == 0 for v in values if v is not None) and any(v is not None for v in values)


def process_gear_batch(
    master_gear_file: Any,
    source_files: list[Any],
    config: GearConfig,
    dry_run: bool = False,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> GearBatchResult:
    """Iteratively processes source workbooks against a master Gear workbook."""
    config.validate()

    # Spool master workbook to disk
    master_path = spool_uploaded_file(master_gear_file, prefix="gear_master_")
    output_path: Path | None = None
    master_wb = None

    logs: list[dict[str, Any]] = []
    processed_count = 0
    skipped_count = 0
    total_formulas: list[dict[str, str]] = []

    try:
        master_wb = openpyxl.load_workbook(master_path)
        if config.gear_sheet_name not in master_wb.sheetnames:
            available = ", ".join(master_wb.sheetnames)
            raise ValueError(
                f"Gear sheet '{config.gear_sheet_name}' not found. Available sheets: {available}"
            )
        sheet = master_wb[config.gear_sheet_name]

        total_files = len(source_files)

        for idx, src_file in enumerate(source_files):
            filename = getattr(src_file, "name", f"file_{idx + 1}.xlsx")
            if progress_callback:
                progress_callback(idx + 1, total_files, filename)

            src_path = spool_uploaded_file(src_file, prefix=f"gear_src_{idx}_")
            src_wb = None
            try:
                src_wb = openpyxl.load_workbook(src_path, data_only=True)
                visible_sheets = [s for s in src_wb.worksheets if s.sheet_state != "hidden"]
                if not visible_sheets:
                    raise ValueError("Workbook has no visible worksheets.")
                try:
                    src_sheet = visible_sheets[config.source_sheet_index]
                except IndexError:
                    raise ValueError(f"Sheet index {config.source_sheet_index} out of range.")

                tank_rows = extract_tank_data(src_sheet)
                meter_sales = extract_meter_sales(src_sheet, tank_rows)
                all_tanks = sorted(set(tank_rows) | set(meter_sales))

                # Match station block in master workbook
                columns, station_name, score = match_station_columns(sheet, config, filename, set(all_tanks))

                if not columns:
                    skipped_count += 1
                    logs.append({
                        "file": filename,
                        "station": station_name,
                        "tank": "-",
                        "meter": None,
                        "dipping": None,
                        "actual_gear": None,
                        "expected_gear": None,
                        "destination_col": "-",
                        "status": "Skipped",
                        "message": f"No matching PMS/AGO tank columns found (score={score:.2f})",
                    })
                    continue

                for tank in all_tanks:
                    col_idx = columns.get(tank)
                    tank_values = tank_rows.get(tank)
                    meter_value = meter_sales.get(tank)

                    if meter_value is None and tank_values and _is_zero_tank_record(tank_values):
                        meter_value = Decimal("0")

                    if tank_values and tank_values.joined_or_empty:
                        logs.append({
                            "file": filename,
                            "station": station_name,
                            "tank": tank,
                            "meter": float(meter_value) if meter_value is not None else None,
                            "dipping": None,
                            "actual_gear": None,
                            "expected_gear": None,
                            "destination_col": get_column_letter(col_idx) if col_idx else "-",
                            "status": "Skipped",
                            "message": "Joined / empty tank row",
                        })
                        continue

                    if col_idx is None:
                        logs.append({
                            "file": filename,
                            "station": station_name,
                            "tank": tank,
                            "meter": float(meter_value) if meter_value is not None else None,
                            "dipping": float(tank_values.dipping_sales) if tank_values and tank_values.dipping_sales is not None else None,
                            "actual_gear": float(tank_values.actual_gear) if tank_values and tank_values.actual_gear is not None else None,
                            "expected_gear": float(tank_values.expected_gear) if tank_values and tank_values.expected_gear is not None else None,
                            "destination_col": "-",
                            "status": "Skipped",
                            "message": "No matching column for tank in station block",
                        })
                        continue

                    if is_protected_column(sheet, col_idx, config.protected_last_columns_count):
                        logs.append({
                            "file": filename,
                            "station": station_name,
                            "tank": tank,
                            "meter": float(meter_value) if meter_value is not None else None,
                            "dipping": None,
                            "actual_gear": None,
                            "expected_gear": None,
                            "destination_col": get_column_letter(col_idx),
                            "status": "Skipped",
                            "message": "Target column is protected",
                        })
                        continue

                    writes = [
                        (config.sales_meter_row, meter_value),
                        (config.sales_dipping_row, tank_values.dipping_sales if tank_values else None),
                        (config.actual_gear_row, tank_values.actual_gear if tank_values else None),
                        (config.expected_gear_row, tank_values.expected_gear if tank_values else None),
                    ]

                    if not dry_run:
                        for row_idx, val in writes:
                            if val is not None:
                                sheet.cell(row_idx, col_idx).value = decimal_to_excel(val)

                        col_letter = get_column_letter(col_idx)
                        sheet.cell(config.gear_percent_row, col_idx).value = (
                            f"={col_letter}{config.actual_gear_row}/{col_letter}{config.expected_gear_row}*100"
                        )

                        prev_formula = sheet.cell(config.previous_average_gear_row, col_idx).value
                        avg_formula = build_average_formula(prev_formula, col_idx, config.gear_percent_row, meter_value)
                        sheet.cell(config.average_gear_row, col_idx).value = avg_formula

                    logs.append({
                        "file": filename,
                        "station": station_name,
                        "tank": tank,
                        "meter": float(meter_value) if meter_value is not None else None,
                        "dipping": float(tank_values.dipping_sales) if tank_values and tank_values.dipping_sales is not None else None,
                        "actual_gear": float(tank_values.actual_gear) if tank_values and tank_values.actual_gear is not None else None,
                        "expected_gear": float(tank_values.expected_gear) if tank_values and tank_values.expected_gear is not None else None,
                        "destination_col": get_column_letter(col_idx),
                        "status": "Preview" if dry_run else "Written",
                        "message": f"OK (score={score:.2f})",
                    })

                processed_count += 1
            except Exception as exc:
                skipped_count += 1
                logs.append({
                    "file": filename,
                    "station": "-",
                    "tank": "-",
                    "meter": None,
                    "dipping": None,
                    "actual_gear": None,
                    "expected_gear": None,
                    "destination_col": "-",
                    "status": "Error",
                    "message": str(exc),
                })
            finally:
                if src_wb:
                    try:
                        src_wb.close()
                    except Exception:
                        pass
                cleanup_files(src_path)

        if not dry_run:
            total_formulas = build_total_formulas(sheet, config)
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", prefix="GEAR_FILLED_") as out_tmp:
                output_path = Path(out_tmp.name)
            master_wb.save(output_path)
            master_wb.close()
            master_wb = None

            # Restore enhanced conditional formatting (green/yellow/red colors) stripped by openpyxl
            preserve_gear_extensions(
                source_master_path=master_path,
                generated_output_path=output_path,
                gear_sheet_name=config.gear_sheet_name,
                gear_percent_row=config.gear_percent_row,
                average_gear_row=config.average_gear_row,
            )

    finally:
        if master_wb:
            try:
                master_wb.close()
            except Exception:
                pass
        cleanup_files(master_path)
        force_gc()

    return GearBatchResult(
        output_path=output_path,
        processed_count=processed_count,
        skipped_count=skipped_count,
        logs=logs,
        total_formulas_written=total_formulas,
    )
