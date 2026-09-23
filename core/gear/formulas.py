from __future__ import annotations

from decimal import Decimal
import re

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import GearConfig
from .utils import normalize_tank_name


DIVISOR_PATTERN = re.compile(r"/\s*(\d+(?:\.\d+)?)\s*\)?\s*$")
CELL_REF_PATTERN = re.compile(r"\$?[A-Z]{1,3}\$?\d+")


def extract_divisor(formula: str | None) -> int:
    if not formula or not isinstance(formula, str):
        return 0
    match = DIVISOR_PATTERN.search(formula)
    if not match:
        return 0
    return int(Decimal(match.group(1)))


def extract_formula_refs(formula: str | None) -> list[str]:
    if not formula or not isinstance(formula, str):
        return []
    return CELL_REF_PATTERN.findall(formula)


def build_average_formula(
    previous_formula: str | None,
    column_index: int,
    today_gear_percent_row: int,
    meter_sales: Decimal | None,
) -> str:
    refs = extract_formula_refs(previous_formula)
    today_ref = f"{get_column_letter(column_index)}{today_gear_percent_row}"
    if today_ref not in refs:
        refs.append(today_ref)
    divisor = extract_divisor(previous_formula)
    if meter_sales is not None and meter_sales > 0:
        divisor += 1
    if divisor <= 0:
        divisor = 1 if meter_sales is not None and meter_sales > 0 else max(len(refs), 1)
    return f"=({'+'.join(refs)})/{divisor}"


def build_total_formulas(sheet: Worksheet, config: GearConfig) -> list[dict[str, str]]:
    """Generates and writes total column formulas for PMS and AGO."""
    total_columns = _find_total_columns(sheet, config)
    written = []

    for col_idx in total_columns:
        col_letter = get_column_letter(col_idx)
        fuel = str(sheet.cell(config.header_row_for_tanks, col_idx).value or "").strip().upper()

        meter_formula = _total_sum_formula(sheet, config, fuel, config.sales_meter_row)
        dipping_formula = _total_sum_formula(sheet, config, fuel, config.sales_dipping_row)
        if meter_formula:
            sheet.cell(config.sales_meter_row, col_idx).value = meter_formula
        if dipping_formula:
            sheet.cell(config.sales_dipping_row, col_idx).value = dipping_formula

        sheet.cell(config.actual_gear_row, col_idx).value = (
            f"={col_letter}{config.sales_meter_row}-{col_letter}{config.sales_dipping_row}"
        )
        sheet.cell(config.expected_gear_row, col_idx).value = (
            f"={col_letter}{config.sales_meter_row}*0.085"
        )
        sheet.cell(config.gear_percent_row, col_idx).value = (
            f"={col_letter}{config.actual_gear_row}/{col_letter}{config.expected_gear_row}*100"
        )

        prev_avg = sheet.cell(config.previous_average_gear_row, col_idx).value
        average_formula = _build_total_average_formula(prev_avg, col_idx, config.gear_percent_row)
        sheet.cell(config.average_gear_row, col_idx).value = average_formula

        written.append({
            "column": col_letter,
            "fuel": fuel,
            "meter_formula": meter_formula or "",
            "average_formula": average_formula,
        })

    return written


def _find_total_columns(sheet: Worksheet, config: GearConfig) -> list[int]:
    columns: list[int] = []
    station_row = config.header_row_for_tanks - 1
    for col_idx in range(1, sheet.max_column + 1):
        station = str(sheet.cell(station_row, col_idx).value or "").strip().upper()
        header = str(sheet.cell(config.header_row_for_tanks, col_idx).value or "").strip().upper()
        if station == "TOTAL" and header in {"PMS", "AGO"}:
            columns.append(col_idx)
    return columns


def _total_sum_formula(sheet: Worksheet, config: GearConfig, fuel: str, row_idx: int) -> str | None:
    refs: list[str] = []
    station_row = config.header_row_for_tanks - 1
    for col_idx in range(1, sheet.max_column + 1):
        station = str(sheet.cell(station_row, col_idx).value or "").strip().upper()
        if station == "TOTAL":
            continue
        tank = normalize_tank_name(sheet.cell(config.header_row_for_tanks, col_idx).value)
        if tank and tank.startswith(fuel):
            refs.append(f"{get_column_letter(col_idx)}{row_idx}")
    if not refs:
        return None
    return "=" + "+".join(refs)


def _build_total_average_formula(previous_formula: object, col_idx: int, gear_percent_row: int) -> str:
    previous_text = previous_formula if isinstance(previous_formula, str) else None
    refs = extract_formula_refs(previous_text)
    today_ref = f"{get_column_letter(col_idx)}{gear_percent_row}"
    if today_ref not in refs:
        refs.append(today_ref)
    divisor = extract_divisor(previous_text) + 1
    if divisor <= 0:
        divisor = max(len(refs), 1)
    return f"=({'+'.join(refs)})/{divisor}"
