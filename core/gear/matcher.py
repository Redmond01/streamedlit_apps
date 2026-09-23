from __future__ import annotations

from difflib import SequenceMatcher
from pathlib import Path
import re
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .config import GearConfig
from .utils import normalize_tank_name


def is_protected_column(sheet: Worksheet, col_idx: int, protected_last_columns_count: int) -> bool:
    first_protected = sheet.max_column - protected_last_columns_count + 1
    return protected_last_columns_count > 0 and col_idx >= first_protected


def match_station_columns(
    sheet: Worksheet,
    config: GearConfig,
    source_filename: str,
    source_tanks: set[str],
) -> tuple[dict[str, int], str, float]:
    """Matches a source file to a station column block in the master Gear sheet.
    
    Returns:
        (tank_to_column_dict, matched_station_name, match_score)
    """
    block = _station_block_for_source(sheet, config, source_filename)
    if block:
        start_col, end_col, station_name, score = block
    else:
        start_col, end_col = 1, sheet.max_column
        station_name = "UNMATCHED (FULL SCAN)"
        score = 0.0

    block_tanks: list[tuple[int, str]] = []
    for col_idx in range(start_col, end_col + 1):
        if is_protected_column(sheet, col_idx, config.protected_last_columns_count):
            continue
        tank = normalize_tank_name(sheet.cell(config.header_row_for_tanks, col_idx).value)
        if tank and (tank.startswith("PMS") or tank.startswith("AGO")):
            block_tanks.append((col_idx, tank))

    resolved = _resolve_tank_columns(block_tanks, source_tanks)
    return resolved, station_name, score


def _station_block_for_source(
    sheet: Worksheet,
    config: GearConfig,
    source_filename: str,
) -> tuple[int, int, str, float] | None:
    station_row = config.header_row_for_tanks - 1
    if station_row < 1:
        return None
    source_key = _normalize_station_name(Path(source_filename).stem)
    candidates = _station_blocks(sheet, station_row)
    if not candidates:
        return None

    scored: list[tuple[float, int, int, str]] = []
    for start_col, end_col, station_name in candidates:
        station_key = _normalize_station_name(station_name)
        score = _station_match_score(source_key, station_key)
        scored.append((score, start_col, end_col, station_name))
    scored.sort(reverse=True, key=lambda item: item[0])
    best_score, start_col, end_col, station_name = scored[0]
    if best_score < 0.45:
        return None
    return start_col, end_col, station_name, best_score


def _station_blocks(sheet: Worksheet, station_row: int) -> list[tuple[int, int, str]]:
    blocks: list[tuple[int, int, str]] = []
    for merged_range in sheet.merged_cells.ranges:
        if merged_range.min_row <= station_row <= merged_range.max_row:
            value = sheet.cell(merged_range.min_row, merged_range.min_col).value
            if value and str(value).strip().upper() not in {"STATION", "TOTAL"}:
                blocks.append((merged_range.min_col, merged_range.max_col, str(value).strip()))

    occupied = {(start, end) for start, end, _ in blocks}
    non_empty_cols = [
        col_idx
        for col_idx in range(1, sheet.max_column + 1)
        if sheet.cell(station_row, col_idx).value not in (None, "")
    ]
    for idx, col_idx in enumerate(non_empty_cols):
        value = str(sheet.cell(station_row, col_idx).value).strip()
        if not value or value.upper() in {"STATION", "TOTAL"}:
            continue
        if any(start <= col_idx <= end for start, end in occupied):
            continue
        next_col = non_empty_cols[idx + 1] if idx + 1 < len(non_empty_cols) else sheet.max_column + 1
        blocks.append((col_idx, next_col - 1, value))
    return sorted(blocks, key=lambda item: item[0])


def _normalize_station_name(value: str) -> str:
    text = value.upper()
    text = re.sub(r"\.XLSX?$|\.XLSM$", " ", text)
    text = re.sub(r"\(([^)]*)\)", _parenthetical_station_text, text)
    text = re.sub(r"\bJUNE\b|\bJULY\b|\bAUGUST\b|\bSEPTEMBER\b", " ", text)
    text = text.replace("TOLL GATE", "TOLLGATE")
    text = re.sub(r"\b(SGR|OBX|TA|AGS)\b", " ", text)
    text = re.sub(r"\bGATE\b", " ", text)
    text = re.sub(r"\bODE\b", " ", text)
    text = re.sub(r"\bROAD\b", " ", text)
    text = text.replace("OGBOMOSHO", "OGBOMOSO")
    text = text.replace("IJEBUODE", "IJEBU")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _parenthetical_station_text(match: re.Match[str]) -> str:
    inner = match.group(1).strip()
    if re.search(r"[A-Z]", inner):
        return f" {inner} "
    return " "


def _station_match_score(source_key: str, station_key: str) -> float:
    if not source_key or not station_key:
        return 0.0
    source_tokens = set(source_key.split())
    station_tokens = set(station_key.split())
    overlap = len(source_tokens & station_tokens) / max(len(source_tokens | station_tokens), 1)
    ratio = SequenceMatcher(None, source_key, station_key).ratio()
    substring_bonus = 0.25 if source_key in station_key or station_key in source_key else 0.0
    first_token_bonus = 0.25 if (source_tokens and station_tokens and source_key.split()[0] == station_key.split()[0]) else 0.0
    return max(overlap, ratio) + substring_bonus + first_token_bonus


def _resolve_tank_columns(block_tanks: list[tuple[int, str]], source_tanks: set[str]) -> dict[str, int]:
    columns: dict[str, int] = {}
    used_cols: set[int] = set()

    for fuel in ("PMS", "AGO"):
        fuel_columns = [(col, tank) for col, tank in block_tanks if tank.startswith(fuel)]
        source_fuel_tanks = sorted(
            [tank for tank in source_tanks if tank.startswith(fuel)],
            key=lambda value: int(re.search(r"\d+", value).group(0)),
        )
        header_names = [tank for _, tank in fuel_columns]
        has_duplicate_headers = len(header_names) != len(set(header_names))
        missing_source_headers = any(tank not in header_names for tank in source_fuel_tanks)

        if source_fuel_tanks and len(fuel_columns) == len(source_fuel_tanks) and (has_duplicate_headers or missing_source_headers):
            for (col_idx, _), tank in zip(fuel_columns, source_fuel_tanks):
                columns[tank] = col_idx
                used_cols.add(col_idx)

    for col_idx, tank in block_tanks:
        if col_idx in used_cols or tank in columns:
            continue
        columns[tank] = col_idx

    unresolved_source_tanks = [tank for tank in _sort_tanks(source_tanks) if tank not in columns]
    if unresolved_source_tanks and len(block_tanks) == len(source_tanks):
        return {tank: col_idx for (col_idx, _), tank in zip(block_tanks, _sort_tanks(source_tanks))}
    return columns


def _sort_tanks(tanks: set[str]) -> list[str]:
    fuel_order = {"PMS": 0, "AGO": 1}
    def key(value: str) -> tuple[int, int, str]:
        match = re.match(r"([A-Z]+)(\d+)", value)
        if not match:
            return (99, 999, value)
        return (fuel_order.get(match.group(1), 99), int(match.group(2)), value)

    return sorted(tanks, key=key)
