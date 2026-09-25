from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
import logging
from pathlib import Path
import tempfile
from typing import Any, Callable

import openpyxl
from openpyxl.utils import get_column_letter

from core.gear.utils import decimal_to_excel
from core.memory import cleanup_files, force_gc, spool_uploaded_file
from .config import SalesConfig
from .extractors import extract_station_sales_records
from .matcher import build_date_row_map, build_station_column_map, match_station_column

logger = logging.getLogger(__name__)


@dataclass
class SalesBatchResult:
    output_path: Path | None
    processed_count: int
    skipped_count: int
    total_records_written: int
    logs: list[dict[str, Any]]
    summary_metrics: dict[str, Any]


def process_sales_batch(
    master_sales_file: Any,
    source_files: list[Any],
    config: SalesConfig,
    dry_run: bool = False,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> SalesBatchResult:
    """Iteratively processes daily station workbooks and aggregates sales into master sales report."""
    config.validate()

    master_path = spool_uploaded_file(master_sales_file, prefix="sales_master_")
    output_path: Path | None = None
    master_wb = None

    logs: list[dict[str, Any]] = []
    processed_count = 0
    skipped_count = 0
    total_records_written = 0

    total_pms_litres = Decimal("0")
    total_ago_litres = Decimal("0")
    total_lpg_kg = Decimal("0")

    try:
        master_wb = openpyxl.load_workbook(master_path, data_only=False)

        # Validate master sheets
        for sheet_name in (config.pms_sheet_name, config.ago_sheet_name):
            if sheet_name not in master_wb.sheetnames:
                available = ", ".join(master_wb.sheetnames)
                raise ValueError(
                    f"Required sheet '{sheet_name}' not found in master workbook. Available sheets: {available}"
                )

        pms_ws = master_wb[config.pms_sheet_name]
        ago_ws = master_wb[config.ago_sheet_name]
        lpg_ws = master_wb[config.lpg_sheet_name] if config.lpg_sheet_name in master_wb.sheetnames else None

        # Build column maps for stations
        pms_cols = build_station_column_map(pms_ws, config.pms_header_row)
        ago_cols = build_station_column_map(ago_ws, config.ago_header_row)
        lpg_cols = build_station_column_map(lpg_ws, config.lpg_header_row) if lpg_ws else {}

        # Build date to row lookup maps
        pms_dates = build_date_row_map(pms_ws)
        ago_dates = build_date_row_map(ago_ws)
        lpg_dates = build_date_row_map(lpg_ws) if lpg_ws else {}

        total_files = len(source_files)

        for idx, src_file in enumerate(source_files):
            filename = getattr(src_file, "name", f"station_{idx + 1}.xlsx")
            if progress_callback:
                progress_callback(idx + 1, total_files, filename)

            src_path = spool_uploaded_file(src_file, prefix=f"sales_src_{idx}_")
            src_wb = None
            try:
                src_wb = openpyxl.load_workbook(src_path, data_only=True)
                records = extract_station_sales_records(src_wb, filename, config)

                if not records:
                    skipped_count += 1
                    logs.append({
                        "file": filename,
                        "station": "-",
                        "date": "-",
                        "sheet": "-",
                        "pms_sales": None,
                        "ago_sales": None,
                        "lpg_sales": None,
                        "target_cols": "-",
                        "status": "Skipped",
                        "message": "No valid daily sales records found in workbook",
                    })
                    continue

                # Match station columns
                pms_col, pms_station, pms_score = match_station_column(filename, pms_cols, config.match_threshold)
                ago_col, ago_station, ago_score = match_station_column(filename, ago_cols, config.match_threshold)
                lpg_col, lpg_station, lpg_score = match_station_column(filename, lpg_cols, config.match_threshold)

                matched_station_name = pms_station if pms_col else (ago_station if ago_col else lpg_station)

                if pms_col is None and ago_col is None and lpg_col is None:
                    skipped_count += 1
                    logs.append({
                        "file": filename,
                        "station": matched_station_name,
                        "date": "-",
                        "sheet": "-",
                        "pms_sales": None,
                        "ago_sales": None,
                        "lpg_sales": None,
                        "target_cols": "-",
                        "status": "Skipped",
                        "message": f"Could not match station to PMS/AGO/LPG columns (score={max(pms_score, ago_score):.2f})",
                    })
                    continue

                # Process each day record for this station
                for rec in records:
                    date_str = rec.date.strftime("%Y-%m-%d")
                    target_cols_desc: list[str] = []
                    status_desc = "Preview" if dry_run else "Written"
                    notes: list[str] = []

                    # 1. PMS Sales Write
                    if rec.pms is not None:
                        total_pms_litres += rec.pms
                        if pms_col is not None:
                            row_idx = pms_dates.get(rec.date)
                            if row_idx is not None:
                                if not dry_run:
                                    pms_ws.cell(row_idx, pms_col).value = decimal_to_excel(rec.pms)
                                total_records_written += 1
                                target_cols_desc.append(f"PMS:{get_column_letter(pms_col)}{row_idx}")
                            else:
                                notes.append(f"PMS date {date_str} not in master rows")
                        else:
                            notes.append("No PMS column matched")

                    # 2. AGO Sales Write
                    if rec.ago is not None:
                        total_ago_litres += rec.ago
                        if ago_col is not None:
                            row_idx = ago_dates.get(rec.date)
                            if row_idx is not None:
                                if not dry_run:
                                    ago_ws.cell(row_idx, ago_col).value = decimal_to_excel(rec.ago)
                                total_records_written += 1
                                target_cols_desc.append(f"AGO:{get_column_letter(ago_col)}{row_idx}")
                            else:
                                notes.append(f"AGO date {date_str} not in master rows")
                        else:
                            notes.append("No AGO column matched")

                    # 3. LPG Sales Write
                    if rec.lpg is not None:
                        total_lpg_kg += rec.lpg
                        if lpg_col is not None and lpg_ws is not None:
                            row_idx = lpg_dates.get(rec.date)
                            if row_idx is not None:
                                if not dry_run:
                                    lpg_ws.cell(row_idx, lpg_col).value = decimal_to_excel(rec.lpg)
                                total_records_written += 1
                                target_cols_desc.append(f"LPG:{get_column_letter(lpg_col)}{row_idx}")
                            else:
                                notes.append(f"LPG date {date_str} not in master rows")
                        else:
                            notes.append("No LPG column matched")

                    msg = "; ".join(notes) if notes else f"Matched (score={pms_score:.2f})"
                    logs.append({
                        "file": filename,
                        "station": matched_station_name,
                        "date": date_str,
                        "sheet": rec.sheet_name,
                        "pms_sales": float(rec.pms) if rec.pms is not None else None,
                        "ago_sales": float(rec.ago) if rec.ago is not None else None,
                        "lpg_sales": float(rec.lpg) if rec.lpg is not None else None,
                        "target_cols": ", ".join(target_cols_desc) if target_cols_desc else "-",
                        "status": status_desc,
                        "message": msg,
                    })

                processed_count += 1

            except Exception as exc:
                skipped_count += 1
                logger.exception("Error processing file %s", filename)
                logs.append({
                    "file": filename,
                    "station": "-",
                    "date": "-",
                    "sheet": "-",
                    "pms_sales": None,
                    "ago_sales": None,
                    "lpg_sales": None,
                    "target_cols": "-",
                    "status": "Error",
                    "message": str(exc),
                })
            finally:
                if src_wb:
                    try:
                        src_wb.close()
                    except Exception:
                        pass
                cleanup_files(src_path)

        if not dry_run:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx", prefix="SALES_REPORT_") as out_tmp:
                output_path = Path(out_tmp.name)
            master_wb.save(output_path)
            master_wb.close()
            master_wb = None

    finally:
        if master_wb:
            try:
                master_wb.close()
            except Exception:
                pass
        cleanup_files(master_path)
        force_gc()

    return SalesBatchResult(
        output_path=output_path,
        processed_count=processed_count,
        skipped_count=skipped_count,
        total_records_written=total_records_written,
        logs=logs,
        summary_metrics={
            "total_pms_litres": float(total_pms_litres),
            "total_ago_litres": float(total_ago_litres),
            "total_lpg_kg": float(total_lpg_kg),
        },
    )
