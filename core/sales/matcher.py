from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import re
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from core.gear.matcher import _normalize_station_name, _station_match_score
from .extractors import _parse_date_str


def build_station_column_map(sheet: Worksheet, header_row: int) -> dict[int, str]:
    """Builds a mapping from column index to station header name.
    
    Excludes Column 1 (DATE) and the summary TOTAL column.
    """
    column_map: dict[int, str] = {}
    max_c = sheet.max_column or 70
    for col_idx in range(2, max_c + 1):
        val = sheet.cell(header_row, col_idx).value
        if val is None:
            continue
        cleaned = str(val).strip()
        if not cleaned or cleaned.upper() in {"TOTAL", "DATE", "STATION"}:
            continue
        column_map[col_idx] = cleaned
    return column_map


def match_station_column(
    station_source_name: str,
    column_map: dict[int, str],
    threshold: float = 0.50,
) -> tuple[int | None, str, float]:
    """Matches a station source file or name to a column in the destination sheet.
    
    Returns:
        (col_index, matched_station_name, score)
    """
    if not column_map:
        return None, "NO_COLUMNS", 0.0

    source_stem = Path(station_source_name).stem
    source_key = _normalize_station_name(source_stem)

    candidates: list[tuple[float, int, str]] = []
    for col_idx, header_name in column_map.items():
        station_key = _normalize_station_name(header_name)
        score = _station_match_score(source_key, station_key)
        candidates.append((score, col_idx, header_name))

    candidates.sort(key=lambda item: item[0], reverse=True)
    best_score, best_col, best_name = candidates[0]

    if best_score < threshold:
        return None, f"UNMATCHED ({best_name})", best_score

    return best_col, best_name, best_score


def build_date_row_map(sheet: Worksheet) -> dict[date, int]:
    """Scans Column A of a sheet to build a mapping from calendar date to row index."""
    date_map: dict[date, int] = {}
    max_r = sheet.max_row or 40

    for row_idx in range(1, max_r + 1):
        val = sheet.cell(row_idx, 1).value
        if isinstance(val, (datetime, date)):
            d = val.date() if isinstance(val, datetime) else val
            date_map[d] = row_idx
        elif isinstance(val, str):
            parsed = _parse_date_str(val)
            if parsed:
                date_map[parsed] = row_idx

    return date_map
