from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
import logging
import re
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

from core.gear.utils import is_blank, parse_number
from .config import SalesConfig

logger = logging.getLogger(__name__)

DATE_REGEX = re.compile(r"(\d{4}[-/]\d{1,2}[-/]\d{1,2})|(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})")


@dataclass(frozen=True)
class DailySalesRecord:
    station_file: str
    sheet_name: str
    date: date
    pms: Decimal | None
    ago: Decimal | None
    lpg: Decimal | None


def extract_date_from_sheet(sheet: Worksheet) -> date | None:
    """Extracts calendar date from daily station worksheet."""
    max_r = min(sheet.max_row or 10, 8)
    max_c = min(sheet.max_column or 15, 12)

    # Pass 1: Direct datetime/date objects in top rows
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            val = sheet.cell(r, c).value
            if isinstance(val, (datetime, date)):
                return val.date() if isinstance(val, datetime) else val

    # Pass 2: Look for 'DATE' label
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            val = str(sheet.cell(r, c).value or "").strip().upper()
            if val in {"DATE", "DATE:", "TRANS DATE", "TRANS DATE:"}:
                for check_cell in [
                    sheet.cell(r, c + 1).value,
                    sheet.cell(r + 1, c).value,
                    sheet.cell(r, c + 2).value,
                ]:
                    if isinstance(check_cell, (datetime, date)):
                        return check_cell.date() if isinstance(check_cell, datetime) else check_cell
                    if isinstance(check_cell, str):
                        parsed = _parse_date_str(check_cell)
                        if parsed:
                            return parsed

    # Pass 3: Regex string matching across top cells
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            val = sheet.cell(r, c).value
            if isinstance(val, str):
                parsed = _parse_date_str(val)
                if parsed:
                    return parsed

    # Pass 4: Fallback to sheet title
    return _parse_date_str(sheet.title)


def _parse_date_str(text: str) -> date | None:
    if not text:
        return None
    match = DATE_REGEX.search(text.strip())
    if not match:
        return None
    cleaned = match.group(0).replace("/", "-")
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d-%m-%y", "%m-%d-%Y"):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    return None


def extract_sales_from_sheet(sheet: Worksheet) -> dict[str, Decimal]:
    """Extracts daily PMS, AGO, and LPG sales volumes from summary block."""
    sales: dict[str, Decimal] = {}
    max_r = min(sheet.max_row or 30, 25)
    max_c = min(sheet.max_column or 35, 30)

    # Strategy: Scan for 'SALES' header column first for highest precision
    sales_header_col = None
    sales_header_row = None
    for r in range(1, max_r + 1):
        for c in range(10, max_c + 1):
            val = str(sheet.cell(r, c).value or "").strip().upper()
            if val == "SALES":
                # Check if rows directly beneath have PMS/AGO
                below_vals = {
                    str(sheet.cell(r + offset, c).value or "").strip().upper()
                    for offset in range(1, 4)
                }
                if below_vals & {"PMS", "AGO", "LPG"}:
                    sales_header_col = c
                    sales_header_row = r
                    break
        if sales_header_col:
            break

    if sales_header_col and sales_header_row:
        # Extract from the structured sales column
        for r in range(sales_header_row + 1, min(sales_header_row + 8, max_r + 1)):
            label = str(sheet.cell(r, sales_header_col).value or "").strip().upper()
            if label in {"PMS", "AGO", "LPG"} and label not in sales:
                num = parse_number(sheet.cell(r, sales_header_col + 1).value)
                if num is not None:
                    sales[label] = num
            elif label.startswith("TOTAL") or label in {"EXPENSES", "PRODUCTS RECEIVED"}:
                break

    # Secondary scan: If any product is missing, scan general grid for exact tokens
    for product in ("PMS", "AGO", "LPG"):
        if product in sales:
            continue
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                val = sheet.cell(r, c).value
                if isinstance(val, str) and val.strip().upper() == product:
                    # Check cell to right
                    num = parse_number(sheet.cell(r, c + 1).value)
                    if num is not None:
                        sales[product] = num
                        break
            if product in sales:
                break

    return sales


def extract_station_sales_records(
    workbook: Workbook,
    filename: str,
    config: SalesConfig,
) -> list[DailySalesRecord]:
    """Extracts daily sales records from a station workbook based on config sync mode."""
    visible_sheets = [s for s in workbook.worksheets if s.sheet_state != "hidden"]
    if not visible_sheets:
        return []

    target_sheets: list[Worksheet] = []

    if config.sync_mode == "latest":
        try:
            target_sheets = [visible_sheets[config.source_sheet_index]]
        except IndexError:
            target_sheets = [visible_sheets[-1]]
    elif config.sync_mode == "specific_date":
        for s in visible_sheets:
            d = extract_date_from_sheet(s)
            if d and d == config.target_date:
                target_sheets.append(s)
        if not target_sheets:
            # Try latest sheet as fallback or return empty
            return []
    else:  # "all"
        target_sheets = visible_sheets

    records: list[DailySalesRecord] = []
    for ws in target_sheets:
        d = extract_date_from_sheet(ws)
        if not d:
            logger.warning("Could not resolve date for sheet %s in %s", ws.title, filename)
            continue

        if config.sync_mode == "specific_date" and d != config.target_date:
            continue

        sales_dict = extract_sales_from_sheet(ws)
        records.append(
            DailySalesRecord(
                station_file=filename,
                sheet_name=ws.title,
                date=d,
                pms=sales_dict.get("PMS"),
                ago=sales_dict.get("AGO"),
                lpg=sales_dict.get("LPG"),
            )
        )

    return records
