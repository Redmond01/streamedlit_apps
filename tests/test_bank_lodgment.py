from __future__ import annotations

from decimal import Decimal
from pathlib import Path
import unittest
import openpyxl

from core.bank_lodgment.config import BankLodgmentConfig, ConfigError
from core.bank_lodgment.engine import process_bank_lodgment
from core.bank_lodgment.matcher import (
    format_lodgment_time,
    normalize_bank_name,
    resolve_station_for_record,
)
from core.bank_lodgment.parser import (
    BankLodgmentRecord,
    normalize_iso_to_sheet_date,
    parse_lodgment_json,
)


class TestBankLodgmentConfig(unittest.TestCase):
    def test_default_config(self):
        cfg = BankLodgmentConfig()
        cfg.validate()
        self.assertIn("+2349115977488", cfg.phone_map)
        self.assertEqual(cfg.phone_map["+2349115977488"], "Cargo 2 OBX")
        self.assertIn("BR_CARGO_1", cfg.branch_id_map)
        self.assertEqual(cfg.branch_id_map["BR_CARGO_1"], "Cargo 1 OBX")


class TestBankLodgmentParser(unittest.TestCase):
    def test_time_formatting(self):
        self.assertEqual(format_lodgment_time("09:46"), "9.46AM")
        self.assertEqual(format_lodgment_time("10:55"), "10.55AM")
        self.assertEqual(format_lodgment_time("12:00"), "12.00PM")
        self.assertEqual(format_lodgment_time("14:36"), "2.36PM")
        self.assertEqual(format_lodgment_time("17:09"), "5.09PM")

    def test_date_normalization(self):
        self.assertEqual(normalize_iso_to_sheet_date("2026-09-25"), "25-09-2026")
        self.assertEqual(normalize_iso_to_sheet_date("25-09-2026"), "25-09-2026")

    def test_bank_normalization(self):
        self.assertEqual(normalize_bank_name("Union Bank"), "UNION")
        self.assertEqual(normalize_bank_name("Access Bank"), "FCMB")
        self.assertEqual(normalize_bank_name("FCMB"), "FCMB")
        self.assertEqual(normalize_bank_name("Zenith Bank"), "ZENITH")

    def test_parse_and_deduplicate(self):
        sample_json = {
            "export_meta": {"target_reconciliation_date": "2026-09-25"},
            "records": [
                {
                    "officer_message_id": "MSG_1",
                    "date": "2026-09-25",
                    "time": "10:00",
                    "bank_officer": {"name": "Lawrence", "bank": "Union Bank"},
                    "confirmation": {"parsed_amount": 500000, "status": "APPROVED"},
                    "branch_submission": {"manager_phone": "+2348000000001", "mapped_branch_id": "BR_TEST"},
                },
                {
                    "officer_message_id": "MSG_2",
                    "date": "2026-09-25",
                    "time": "10:05",
                    "bank_officer": {"name": "Lawrence", "bank": "Union Bank"},
                    "confirmation": {"parsed_amount": 500000, "status": "APPROVED"},
                    "branch_submission": {"manager_phone": "+2348000000001", "mapped_branch_id": "BR_TEST"},
                },
            ],
        }

        meta, recs = parse_lodgment_json(sample_json)
        self.assertEqual(len(recs), 2)
        self.assertFalse(recs[0].is_duplicate)
        self.assertTrue(recs[1].is_duplicate)


class TestBankLodgmentMatcher(unittest.TestCase):
    def setUp(self):
        self.cfg = BankLodgmentConfig()
        self.known_stations = ["Cargo 1 OBX", "Cargo 2 OBX", "Oyo Ilora", "Ife 2 OBX"]

    def test_match_via_branch_id(self):
        rec = BankLodgmentRecord(
            message_id="1", date="2026-09-25", time="10:00",
            bank_officer_name="Abimbola", bank_officer_bank="Access Bank",
            amount=Decimal("100"), status="APPROVED",
            manager_phone="", manager_handle="", raw_caption="Photo",
            mapped_branch_id="BR_CARGO_1", has_receipt_image=True,
        )
        st, strat, score = resolve_station_for_record(rec, self.cfg, self.known_stations)
        self.assertEqual(st, "Cargo 1 OBX")
        self.assertEqual(strat, "branch_id")

    def test_match_via_phone(self):
        rec = BankLodgmentRecord(
            message_id="1", date="2026-09-25", time="10:00",
            bank_officer_name="Lawrence", bank_officer_bank="Union Bank",
            amount=Decimal("100"), status="APPROVED",
            manager_phone="+2349115977488", manager_handle="Olamilekan", raw_caption="Photo",
            mapped_branch_id=None, has_receipt_image=True,
        )
        st, strat, score = resolve_station_for_record(rec, self.cfg, self.known_stations)
        self.assertEqual(st, "Cargo 2 OBX")
        self.assertEqual(strat, "manager_phone")

    def test_match_via_caption(self):
        rec = BankLodgmentRecord(
            message_id="1", date="2026-09-25", time="10:00",
            bank_officer_name="Lawrence", bank_officer_bank="Union Bank",
            amount=Decimal("100"), status="APPROVED",
            manager_phone="", manager_handle="", raw_caption="OYO 1",
            mapped_branch_id=None, has_receipt_image=True,
        )
        st, strat, score = resolve_station_for_record(rec, self.cfg, self.known_stations)
        self.assertEqual(st, "Oyo Ilora")


class TestBankLodgmentIntegration(unittest.TestCase):
    def test_real_files_e2e(self):
        master_file = Path("scratch/SIFEM BANK LODGEMENTS - SEPTEMBER 2026.xlsx")
        json_file = Path("scratch/download.json")

        if not (master_file.exists() and json_file.exists()):
            self.skipTest("Sample scratch workbooks not present.")

        cfg = BankLodgmentConfig()
        result = process_bank_lodgment(
            master_file=master_file,
            json_file=json_file,
            config=cfg,
            dry_run=False,
        )

        self.assertEqual(result.target_sheet_name, "25-09-2026")
        self.assertEqual(result.total_records, 41)
        self.assertEqual(result.confirmed_count, 38)
        self.assertEqual(result.duplicate_count, 3)
        self.assertEqual(result.unmatched_count, 0)
        self.assertEqual(result.total_confirmed_amount, Decimal("58289340"))

        self.assertIsNotNone(result.output_path)
        self.assertTrue(result.output_path.exists())

        # Verify output
        wb = openpyxl.load_workbook(result.output_path, data_only=False)
        ws = wb["25-09-2026"]

        # Row 8: Cargo 2 OBX (800000)
        self.assertEqual(ws.cell(8, 1).value, "Cargo 2 OBX")
        self.assertEqual(ws.cell(8, 3).value, 800000)
        self.assertEqual(ws.cell(8, 7).value, "YES")
        self.assertEqual(ws.cell(8, 8).value, "LAWRENCE")
        self.assertEqual(ws.cell(8, 9).value, "C")

        # Row 9: Cargo 2 OBX (296150)
        self.assertEqual(ws.cell(9, 1).value, "Cargo 2 OBX")
        self.assertEqual(ws.cell(9, 3).value, 296150)
        self.assertEqual(ws.cell(9, 7).value, "YES")
        self.assertEqual(ws.cell(9, 8).value, "LAWRENCE")

        # Verify cell styles and colors (Default: row mode)
        # Confirmed row (Row 8 - Cargo 2 OBX)
        self.assertEqual(ws.cell(8, 7).fill.fill_type, "solid")
        self.assertIn(str(ws.cell(8, 7).fill.start_color.rgb).upper(), ("00C6EFCE", "C6EFCE"))
        self.assertIn(str(ws.cell(8, 7).font.color.rgb).upper(), ("00006100", "006100"))
        self.assertTrue(ws.cell(8, 7).font.bold)
        self.assertIn(str(ws.cell(8, 8).fill.start_color.rgb).upper(), ("00C6EFCE", "C6EFCE"))
        self.assertIn(str(ws.cell(8, 9).fill.start_color.rgb).upper(), ("00C6EFCE", "C6EFCE"))
        self.assertIn(str(ws.cell(8, 1).fill.start_color.rgb).upper(), ("00E2EFDA", "E2EFDA"))

        # Unconfirmed row (Row 2 - Agbara)
        self.assertEqual(ws.cell(2, 7).value, "NO")
        self.assertIn(str(ws.cell(2, 7).fill.start_color.rgb).upper(), ("00FFC7CE", "FFC7CE"))
        self.assertIn(str(ws.cell(2, 7).font.color.rgb).upper(), ("009C0006", "9C0006"))
        self.assertTrue(ws.cell(2, 7).font.bold)
        self.assertIn(str(ws.cell(2, 1).fill.start_color.rgb).upper(), ("00FCE4D6", "FCE4D6"))

        wb.close()
        result.output_path.unlink()

    def test_real_files_columns_color_mode(self):
        master_file = Path("scratch/SIFEM BANK LODGEMENTS - SEPTEMBER 2026.xlsx")
        json_file = Path("scratch/download.json")

        if not (master_file.exists() and json_file.exists()):
            self.skipTest("Sample scratch workbooks not present.")

        cfg = BankLodgmentConfig(color_mode="columns")
        result = process_bank_lodgment(
            master_file=master_file,
            json_file=json_file,
            config=cfg,
            dry_run=False,
        )

        wb = openpyxl.load_workbook(result.output_path, data_only=False)
        ws = wb["25-09-2026"]

        # In columns mode, Cols 7-9 are styled but Col 1 is NOT filled with pastel row accent
        self.assertIn(str(ws.cell(8, 7).fill.start_color.rgb).upper(), ("00C6EFCE", "C6EFCE"))
        self.assertIn(str(ws.cell(8, 8).fill.start_color.rgb).upper(), ("00C6EFCE", "C6EFCE"))
        self.assertIn(str(ws.cell(8, 9).fill.start_color.rgb).upper(), ("00C6EFCE", "C6EFCE"))
        self.assertIn(ws.cell(8, 1).fill.fill_type, (None, "none"))

        wb.close()
        result.output_path.unlink()


if __name__ == "__main__":
    unittest.main()
