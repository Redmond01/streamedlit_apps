from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import unittest
import openpyxl

from core.sales.config import ConfigError, SalesConfig
from core.sales.engine import process_sales_batch
from core.sales.extractors import (
    extract_date_from_sheet,
    extract_sales_from_sheet,
    extract_station_sales_records,
)
from core.sales.matcher import (
    build_date_row_map,
    build_station_column_map,
    match_station_column,
)


class TestSalesConfig(unittest.TestCase):
    def test_default_config_valid(self):
        cfg = SalesConfig()
        cfg.validate()
        self.assertEqual(cfg.pms_sheet_name, "PMS")
        self.assertEqual(cfg.ago_sheet_name, "AGO")
        self.assertEqual(cfg.lpg_sheet_name, "LPG")
        self.assertEqual(cfg.sync_mode, "all")

    def test_invalid_header_row(self):
        cfg = SalesConfig(pms_header_row=0)
        with self.assertRaises(ConfigError):
            cfg.validate()

    def test_invalid_sync_mode(self):
        cfg = SalesConfig(sync_mode="invalid_mode")
        with self.assertRaises(ConfigError):
            cfg.validate()

    def test_specific_date_requires_target_date(self):
        cfg = SalesConfig(sync_mode="specific_date", target_date=None)
        with self.assertRaises(ConfigError):
            cfg.validate()

        cfg_ok = SalesConfig(sync_mode="specific_date", target_date=date(2026, 9, 1))
        cfg_ok.validate()


class TestSalesExtractors(unittest.TestCase):
    def test_extract_date_and_sales_synthetic(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Day1"

        # Date in row 1
        ws["B1"] = datetime(2026, 9, 15, 0, 0)

        # Sales summary table
        ws["U3"] = "SALES"
        ws["U4"] = "PMS"
        ws["V4"] = 12500.50
        ws["U5"] = "AGO"
        ws["V5"] = 8300.25
        ws["U6"] = "LPG"
        ws["V6"] = 450.00

        d = extract_date_from_sheet(ws)
        self.assertEqual(d, date(2026, 9, 15))

        sales = extract_sales_from_sheet(ws)
        self.assertEqual(sales.get("PMS"), Decimal("12500.50"))
        self.assertEqual(sales.get("AGO"), Decimal("8300.25"))
        self.assertEqual(sales.get("LPG"), Decimal("450.00"))


class TestSalesMatcher(unittest.TestCase):
    def test_station_matching(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A2"] = "DATE"
        ws["B2"] = "AGBARA 1"
        ws["C2"] = "SAGAM 3"
        ws["D2"] = "SAGAM 4"
        ws["E2"] = "SAAPADE"
        ws["F2"] = "TOTAL"

        col_map = build_station_column_map(ws, header_row=2)
        self.assertEqual(len(col_map), 4)
        self.assertNotIn(1, col_map)
        self.assertNotIn(6, col_map)

        # Match Sagam 3
        col, name, score = match_station_column("SEPTEMBER Sagam 3 sgr.xlsx", col_map)
        self.assertEqual(col, 3)
        self.assertEqual(name, "SAGAM 3")
        self.assertGreaterEqual(score, 1.0)

        # Match Sagam 4
        col, name, score = match_station_column("SEPTEMBER Sagam 4 sgr.xlsx", col_map)
        self.assertEqual(col, 4)
        self.assertEqual(name, "SAGAM 4")

        # Match Saapade daily
        col, name, score = match_station_column("Saapade daily sales.xlsx", col_map)
        self.assertEqual(col, 5)
        self.assertEqual(name, "SAAPADE")

    def test_date_row_mapping(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "DATE"
        ws["A2"] = datetime(2026, 9, 1, 0, 0)
        ws["A3"] = datetime(2026, 9, 2, 0, 0)
        ws["A4"] = "2026-09-03"

        date_map = build_date_row_map(ws)
        self.assertEqual(date_map.get(date(2026, 9, 1)), 2)
        self.assertEqual(date_map.get(date(2026, 9, 2)), 3)
        self.assertEqual(date_map.get(date(2026, 9, 3)), 4)


class TestSalesIntegration(unittest.TestCase):
    def test_integration_sample_workbooks(self):
        master_file = Path("scratch/SEPTEMBER SALES REPORT 2026.xlsx")
        s3_file = Path("scratch/SEPTEMBER Sagam 3 sgr.xlsx")
        s4_file = Path("scratch/SEPTEMBER Sagam 4 sgr.xlsx")

        if not (master_file.exists() and s3_file.exists() and s4_file.exists()):
            self.skipTest("Sample scratch workbooks not present.")

        cfg = SalesConfig(sync_mode="all")
        result = process_sales_batch(
            master_sales_file=master_file,
            source_files=[s3_file, s4_file],
            config=cfg,
            dry_run=False,
        )

        self.assertEqual(result.processed_count, 2)
        self.assertEqual(result.skipped_count, 0)
        self.assertEqual(result.total_records_written, 92)
        self.assertIsNotNone(result.output_path)
        self.assertTrue(result.output_path.exists())

        # Verify output formulas & values
        out_wb = openpyxl.load_workbook(result.output_path, data_only=False)
        pms_ws = out_wb["PMS"]
        ago_ws = out_wb["AGO"]

        # Sagam 3 (col 57) Day 1 (row 3)
        self.assertAlmostEqual(float(pms_ws.cell(3, 57).value), 12626.72, places=2)
        self.assertAlmostEqual(float(ago_ws.cell(3, 57).value), 13586.58, places=2)

        # Sagam 4 (col 58) Day 1 (row 3)
        self.assertAlmostEqual(float(pms_ws.cell(3, 58).value), 5993.29, places=2)
        self.assertAlmostEqual(float(ago_ws.cell(3, 58).value), 987.01, places=2)

        # Check Total Column formula
        self.assertEqual(pms_ws.cell(3, 62).value, "=SUM(B3:BI3)")
        self.assertEqual(ago_ws.cell(3, 62).value, "=SUM(B3:BI3)")

        # Check Total Row formula
        self.assertEqual(pms_ws.cell(34, 57).value, "=SUM(BE3:BE33)")
        self.assertEqual(pms_ws.cell(34, 62).value, "=SUM(BJ3:BJ33)")

        out_wb.close()
        result.output_path.unlink()


if __name__ == "__main__":
    unittest.main()
