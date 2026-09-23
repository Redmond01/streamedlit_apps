from __future__ import annotations

from pathlib import Path
from core.pos_recon.common import normalize_pos, parse_money, cell_date_to_text
from core.pos_recon.cashbook_extractor import extract_cashbook_loop
from core.pos_recon.bank_loader import load_json_to_workbook
from core.memory import cleanup_files


BANK_DIR = Path("/home/im-redmond/Desktop/Desktop/bank_statement/scripts")


def test_pos_normalization():
    assert normalize_pos("SGR 02") == "SGR 2"
    assert normalize_pos("OBX 113") == "OBX 113"
    assert normalize_pos("T.A 1") == "TA 1"
    assert normalize_pos("TA 04") == "TA 4"
    assert normalize_pos("AGS 005") == "AGS 5"
    assert normalize_pos("UNKNOWN 1") == ""


def test_parse_money():
    assert parse_money("₦12,345.67") == 12345.67
    assert parse_money("(500.00)") == -500.00
    assert parse_money("-250.50") == -250.50
    assert parse_money(" 1,000 ") == 1000.0
    assert parse_money("") is None
    assert parse_money(None) is None


def test_cashbook_extractor_with_real_file():
    sample_file = BANK_DIR / "Book22SEPT.xlsx"
    if not sample_file.exists():
        return

    result = extract_cashbook_loop(
        workbook_input=sample_file,
        date="2026-09-22",
        min_number=11,
        include_audit=True,
    )
    assert result["date"] == "2026-09-22"
    assert "groups" in result
    assert "SGR" in result["groups"]
    assert "OBX" in result["groups"]
    assert len(result["groups"]["SGR"]) > 0


def test_bank_loader_with_real_file():
    sample_wb = BANK_DIR / "Book22SEPT.xlsx"
    sample_json = BANK_DIR / "pos_indexeddb_export_day_22.json"
    if not sample_wb.exists() or not sample_json.exists():
        return

    out_path, stats = load_json_to_workbook(
        workbook_input=sample_wb,
        json_input=sample_json,
        date="2026-09-22",
    )
    try:
        assert out_path.exists()
        assert stats["matched"] > 0
        assert stats["okWritten"] > 0
    finally:
        cleanup_files(out_path)


def test_eod_sync_integration():
    import tempfile
    import openpyxl
    from core.pos_recon.eod_extractor import write_eod_to_workbook

    # Create mock EOD source file
    eod_wb = openpyxl.Workbook()
    ews = eod_wb.active
    ews.title = "Daily_POS_2026_09_22"

    headers = ["Timestamp", "Submission ID", "Station", "Report Date", "Full POS ID", "Amount"]
    ews.append(headers)
    ews.append(["2026-09-22 10:00:00", "SUB1", "Station 1", "2026-09-22", "SGR 12", 15000])
    ews.append(["2026-09-22 11:00:00", "SUB2", "Station 2", "2026-09-22", "OBX 113", 25000])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f_eod:
        eod_wb.save(f_eod.name)
        eod_path = Path(f_eod.name)

    sample_wb = BANK_DIR / "Book22SEPT.xlsx"
    if not sample_wb.exists():
        cleanup_files(eod_path)
        return

    try:
        out_path, stats = write_eod_to_workbook(
            source_input=eod_path,
            target_workbook_input=sample_wb,
            date_str="2026-09-22",
        )
        assert out_path.exists()
        assert stats["sourceRows"] == 2
        assert stats["keptRows"] == 2
        assert stats["mappedRows"] > 0

        # Verify audit sheet was written
        wb_check = openpyxl.load_workbook(out_path)
        assert "EOD Extract Audit" in wb_check.sheetnames
        wb_check.close()
        cleanup_files(out_path)

    finally:
        cleanup_files(eod_path)

