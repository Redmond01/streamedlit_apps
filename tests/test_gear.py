from __future__ import annotations

from decimal import Decimal
import openpyxl

from core.gear.config import GearConfig, ConfigError
from core.gear.formulas import extract_divisor, extract_formula_refs, build_average_formula
from core.gear.utils import normalize_tank_name, parse_number, decimal_to_excel
from core.gear.engine import process_gear_batch


def test_gear_config():
    valid = GearConfig()
    valid.validate()

    invalid = GearConfig(sales_meter_row=130, sales_dipping_row=130)
    try:
        invalid.validate()
        assert False, "Should have raised ConfigError for duplicate rows"
    except ConfigError:
        pass


def test_formula_engine():
    prev_formula = "=(E134+E135)/2"
    refs = extract_formula_refs(prev_formula)
    assert refs == ["E134", "E135"]
    assert extract_divisor(prev_formula) == 2

    # Column 5 is E, today row 134
    avg_formula = build_average_formula(prev_formula, column_index=5, today_gear_percent_row=134, meter_sales=Decimal("500"))
    assert avg_formula == "=(E134+E135)/3"

    # With zero meter sales, divisor shouldn't increment
    avg_zero = build_average_formula(prev_formula, column_index=5, today_gear_percent_row=134, meter_sales=Decimal("0"))
    assert avg_zero == "=(E134+E135)/2"


def test_utils_parsing():
    assert normalize_tank_name("PMS 1") == "PMS1"
    assert normalize_tank_name("AGO 02") == "AGO2"
    assert normalize_tank_name("AG0 3") == "AGO3"
    assert normalize_tank_name("Random Text") is None

    assert parse_number(" 1,234.50 ") == Decimal("1234.50")
    assert parse_number("50%") == Decimal("50")
    assert parse_number(None) is None
    assert parse_number("N/A") is None

    assert decimal_to_excel(Decimal("100.00")) == 100
    assert decimal_to_excel(Decimal("100.50")) == 100.5


def test_process_gear_batch_integration(tmp_path_factory=None):
    import tempfile
    from pathlib import Path
    from core.memory import cleanup_files

    # 1. Create a Master Gear workbook
    master_wb = openpyxl.Workbook()
    ws = master_wb.active
    ws.title = "Sheet1"

    # Row 1: Station headers
    ws["B1"] = "STATION ALPHA"
    ws.merge_cells("B1:D1")
    ws["E1"] = "TOTAL"
    ws["F1"] = "TOTAL"

    # Row 2: Tank headers
    ws["B2"] = "PMS1"
    ws["C2"] = "PMS2"
    ws["D2"] = "AGO1"
    ws["E2"] = "PMS"
    ws["F2"] = "AGO"

    # Row 128: Previous averages
    ws["B128"] = "=(B120)/1"
    ws["C128"] = "=(C120)/1"
    ws["D128"] = "=(D120)/1"
    ws["E128"] = "=(E120)/1"
    ws["F128"] = "=(F120)/1"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f_master:
        master_wb.save(f_master.name)
        master_path = Path(f_master.name)

    # 2. Create a Station Source workbook (Station Alpha.xlsx)
    src_wb = openpyxl.Workbook()
    sws = src_wb.active
    sws.title = "Record"

    # Headers at row 2
    sws["A2"] = "TANK"
    sws["B2"] = "SALES"
    sws["C2"] = "ACTUAL GEAR"
    sws["D2"] = "EXPCTD GEAR"

    # Tank values
    sws["A3"] = "PMS 1"
    sws["B3"] = 1500
    sws["C3"] = 12
    sws["D3"] = 14

    sws["A4"] = "PMS 2"
    sws["B4"] = 2500
    sws["C4"] = 20
    sws["D4"] = 22

    sws["A5"] = "AGO 1"
    sws["B5"] = 800
    sws["C5"] = 8
    sws["D5"] = 9

    # Meter sales section
    sws["F2"] = "TOTAL PMS"
    sws["G2"] = 4000
    sws["F3"] = "TOTAL AGO"
    sws["G3"] = 800

    with tempfile.NamedTemporaryFile(suffix=".xlsx", prefix="Station_Alpha_", delete=False) as f_src:
        src_wb.save(f_src.name)
        src_path = Path(f_src.name)

    try:
        config = GearConfig(
            gear_sheet_name="Sheet1",
            header_row_for_tanks=2,
            sales_meter_row=130,
            sales_dipping_row=131,
            actual_gear_row=132,
            expected_gear_row=133,
            gear_percent_row=134,
            average_gear_row=135,
            previous_average_gear_row=128,
            protected_last_columns_count=0,
            source_sheet_index=-1,
        )

        result = process_gear_batch(
            master_gear_file=master_path,
            source_files=[src_path],
            config=config,
            dry_run=False,
        )

        assert result.processed_count == 1
        assert result.skipped_count == 0
        assert result.output_path is not None
        assert result.output_path.exists()

        # Verify output cells in saved workbook
        out_wb = openpyxl.load_workbook(result.output_path)
        out_ws = out_wb["Sheet1"]

        # Station Alpha PMS1 (col B = 2)
        assert out_ws.cell(131, 2).value == 1500  # Dipping
        assert out_ws.cell(132, 2).value == 12    # Actual gear
        assert out_ws.cell(133, 2).value == 14    # Expected gear
        assert out_ws.cell(135, 2).value == "=(B120+B134)/2" # Average formula

        # Total PMS (col E = 5) formula should be written
        assert out_ws.cell(132, 5).value == "=E130-E131"

        out_wb.close()
        cleanup_files(result.output_path)

    finally:
        cleanup_files(master_path, src_path)

